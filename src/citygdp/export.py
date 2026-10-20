"""导出交付物：指标汇总表、来源登记表、数据说明。

版式见 SPEC §8 与 DESIGN §10：
  * 单个工作表 `城市GDP`，4 个指标块纵向堆叠，行 = 指标×年份，列 = 8 城
  * 指标名每行重复、不合并单元格，保证可被 pandas/CSV 直接读取
  * 空值保持空单元格，不写 0、— 或 N/A
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Config


@dataclass(frozen=True)
class MetricSpec:
    key: str
    label: str
    unit: str
    number_format: str
    decimals: int
    scale: float


def metric_specs(cfg: Config) -> tuple[MetricSpec, ...]:
    """指标的单位、小数位、缩放全部来自 config/params.yaml，避免两处维护。"""
    return tuple(
        MetricSpec(
            key=str(item["key"]),
            label=str(item["label"]),
            unit=str(item["unit"]),
            number_format=str(item["format"]),
            decimals=int(item["decimals"]),
            scale=float(item["scale"]),
        )
        for item in cfg.indicators
    )

# 缺失原因（SPEC §4.3 / DR-02、DR-15）
MISSING_REASONS: dict[tuple[str, int], str] = {
    ("NYC", 2000): "BEA 县级 GDP 序列自 2001 年起，无 2000 年数据",
    ("NYC", 2025): "BEA 县级 2025 年数据排期 2026-12-02 发布",
    ("LDN", 2025): "ONS 区域 GDP 最新版本覆盖至 2024 年",
    ("TYO", 2000): "最新版都民経済計算年報仅覆盖 FY2011–FY2023，跨版本修订约 7%，不拼接（DR-15）",
    ("TYO", 2005): "同上（DR-15）",
    ("TYO", 2010): "同上（DR-15）",
    ("TYO", 2024): "都民経済計算最新为令和 5 年度（FY2023）",
    ("TYO", 2025): "都民経済計算最新为令和 5 年度（FY2023）",
}

FOOTNOTES = (
    "注：",
    "1) 新加坡为城市国家，经济首位度恒为 100%。",
    "2) 香港经济首位度以中国（世界银行 CHN，不含港澳台）为分母。",
    "3) 2024 年起深圳数据含深汕特别合作区，面积同步调整，不作历史回溯。",
    "4) 东京为日本年度（4 月–翌年 3 月），按起始年对齐（令和 5 年度 = FY2023）。",
    "5) 空白表示数据不可得，原因见《数据说明》。",
    "6) 2025 年中国三市、新加坡、香港为初步核算数，其余城市尚无 2025 年数据。",
)

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
BLOCK_FILL_A = PatternFill("solid", fgColor="FFFFFF")
BLOCK_FILL_B = PatternFill("solid", fgColor="F2F7FB")


def _derived_index(conn: sqlite3.Connection) -> dict[tuple[str, int, str], float]:
    rows = conn.execute("SELECT city_code, year, metric, value FROM derived").fetchall()
    return {(r["city_code"], r["year"], r["metric"]): r["value"] for r in rows}


def _row_note(year: int) -> str | None:
    if year == 2025:
        return "中国三市、新加坡、香港为初步数"
    if year == 2024:
        return "深圳自 2024 年起含深汕"
    return None


def export_workbook(conn: sqlite3.Connection, cfg: Config, path: Path) -> Path:
    values = _derived_index(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "城市GDP"

    ws["A1"] = "城市 GDP 主要指标（现价美元口径）"
    ws["A1"].font = Font(size=14, bold=True, color="1F4E79")
    specs = metric_specs(cfg)
    units = "｜".join(f"{m.label} {m.unit}" for m in specs)
    ws["A2"] = (
        f"单位：{units}　按世界银行年度平均汇率折算；空白表示数据不可得"
    )
    ws["A2"].font = Font(size=9, italic=True, color="595959")

    header = ["指标", "年份"] + [c.name_cn for c in cfg.cities] + ["备注"]
    for col, text in enumerate(header, start=1):
        cell = ws.cell(row=3, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    row = 4
    for index, metric in enumerate(specs):
        fill = BLOCK_FILL_A if index % 2 == 0 else BLOCK_FILL_B
        for year in cfg.years:
            label = ws.cell(row=row, column=1, value=f"{metric.label}（{metric.unit}）")
            label.fill = fill
            label.border = BORDER
            year_cell = ws.cell(row=row, column=2, value=year)
            year_cell.alignment = Alignment(horizontal="center")
            year_cell.border = BORDER
            year_cell.fill = fill

            for offset, city in enumerate(cfg.cities, start=3):
                value = values.get((city.code, year, metric.key))
                cell = ws.cell(row=row, column=offset)
                cell.border = BORDER
                cell.fill = fill
                if value is None:
                    continue
                # 不做取整：单元格保留全精度，显示格式由 number_format 控制（DD-04）
                cell.value = value * metric.scale
                cell.number_format = metric.number_format
                cell.alignment = Alignment(horizontal="right")

            note = ws.cell(row=row, column=len(cfg.cities) + 3, value=_row_note(year))
            note.fill = fill
            note.border = BORDER
            if note.value:
                note.font = Font(size=9, color="808080")
            row += 1

    row += 1  # 空行
    for line in FOOTNOTES:
        cell = ws.cell(row=row, column=1, value=line)
        cell.font = Font(size=9, color="595959")
        row += 1

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 8
    for offset in range(3, len(cfg.cities) + 3):
        ws.column_dimensions[get_column_letter(offset)].width = 14
    ws.column_dimensions[get_column_letter(len(cfg.cities) + 3)].width = 32
    ws.freeze_panes = "C4"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_registry(conn: sqlite3.Connection, cfg: Config, path: Path) -> Path:
    coverage: dict[str, tuple[int, int, int]] = {}
    for row in conn.execute(
        "SELECT source_id, MIN(year) AS y0, MAX(year) AS y1, COUNT(*) AS n "
        "FROM fact GROUP BY source_id"
    ):
        coverage[row["source_id"]] = (row["y0"], row["y1"], row["n"])

    statuses: dict[str, set[str]] = {}
    for row in conn.execute("SELECT source_id, status FROM fact GROUP BY source_id, status"):
        statuses.setdefault(row["source_id"], set()).add(row["status"])

    wb = Workbook()
    ws = wb.active
    ws.title = "来源登记"
    headers = [
        "来源ID", "城市", "指标", "来源机构", "数据集", "表号", "URL", "版本",
        "数据更新时间", "访问时间", "ETag", "Last-Modified", "SHA256",
        "本地文件", "获取方式", "覆盖年度", "数据状态", "备注",
    ]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows = conn.execute("SELECT * FROM source ORDER BY id").fetchall()
    for r_index, row in enumerate(rows, start=2):
        y0, y1, _n = coverage.get(row["id"], ("", "", 0))
        covered = f"{y0}–{y1}" if y0 != "" else ""
        status = "/".join(sorted(statuses.get(row["id"], set())))
        local = _relative(row["local_path"], cfg.root)
        updated = row["published_at"] or row["http_last_modified"] or ""
        values = [
            row["id"], row["city_code"] or "", row["indicator"], row["organization"],
            row["dataset_title"], row["table_id"] or "", row["url"], row["vintage"] or "",
            updated, row["accessed_at"], row["http_etag"] or "",
            row["http_last_modified"] or "", row["sha256"] or "", local,
            "自动" if row["acquisition"] == "auto" else "手工", covered, status,
            row["notes"] or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r_index, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (5, 18))

    widths = [30, 8, 16, 26, 42, 16, 55, 18, 12, 22, 16, 26, 30, 40, 10, 12, 18, 40]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _relative(path: str | None, root: Path) -> str:
    if not path:
        return ""
    try:
        return str(Path(path).resolve().relative_to(root))
    except ValueError:
        return path


def export_notes(conn: sqlite3.Connection, cfg: Config, path: Path) -> Path:
    now = datetime.now().astimezone().isoformat(timespec="seconds")
    sources = conn.execute("SELECT id, vintage FROM source ORDER BY id").fetchall()
    derived_count = conn.execute("SELECT COUNT(*) FROM derived").fetchone()[0]
    fact_count = conn.execute("SELECT COUNT(*) FROM fact").fetchone()[0]

    missing_lines = [
        f"| {cfg.city(code).name_cn} | {year} | {reason} |"
        for (code, year), reason in sorted(MISSING_REASONS.items())
    ]

    units = {m.key: m.unit for m in metric_specs(cfg)}

    text = f"""# 数据说明

生成时间：{now}
数据版本：`data/raw` 下 {len(sources)} 个来源快照（见《数据来源登记表》）
记录规模：fact {fact_count} 条、派生指标 {derived_count} 条

## 一、口径

- **价格**：现价（current prices），不做平减。
- **汇率**：世界银行年度平均官方汇率（WDI `PA.NUS.FCRF`），1 美元 = F 单位本币。
- **国家 GDP**：世界银行 `NY.GDP.MKTP.CD`（现价美元）。
- **城市边界**：纽约 = Bronx、Kings、New York、Queens、Richmond 五县之和；伦敦 = ITL1 区域；
  东京 = 都全域；新加坡 = 全国；京沪深 = 全市（深圳 2024 年起含深汕特别合作区）；香港 = 全境。

## 二、指标公式

| 指标 | 公式 | 单位 |
|---|---|---|
| 区域GDP | `GDP_本币 / 汇率 ÷ 10⁸` | {units["gdp_usd"]} |
| 经济首位度 | `城市 GDP（美元）÷ 所在国 GDP（美元）× 100` | {units["primacy_pct"]} |
| 人均GDP | `区域GDP（美元）÷ 人口` | {units["gdp_per_capita_usd"]} |
| 地均GDP | `区域GDP（美元）÷ 面积 ÷ 10⁸` | {units["gdp_per_km2_usd"]} |

## 三、人口与面积口径

- 纽约：5 县合计，7 月 1 日人口估计。
- 伦敦：ONS 年中人口估计。
- 东京：总人口（月初人口单纯平均）。
- 新加坡：年中总人口（含非居民）。
- 北京 / 上海 / 深圳：**年末常住人口**，人均 GDP 使用年平均 =（上年末 + 本年末）÷ 2。
- 香港：年中人口。
- 面积取当年值；新加坡、香港填海造地逐年变化；深圳 2024 年起含深汕（+468.3 km²）。

## 四、缺失数据

| 城市 | 年份 | 原因 |
|---|---|---|
{chr(10).join(missing_lines)}

## 五、数据状态

- `final`：官方正式发布或最终核实值。
- `preliminary`：初步核算数（2025 年中国三市、新加坡、香港）。
- `provisional`：唯一一处为香港 2000 年面积，由 2002 年值回填（DR-14）。

## 六、特殊说明

1. 新加坡为城市国家，首位度恒为 100%。
2. 香港首位度以中国（世界银行 CHN，不含港澳台）为分母；若改用含港口径，份额约低 2%。
3. 东京为日本年度，按起始年对齐；2000/2005/2010 因跨版本修订约 7% 未拼接（DR-15）。
4. 香港面积采用「不含水塘等水域」的分区土地面积，与常见的 1,114.57 km² 相差约 2.3%（DR-13）。

## 七、来源清单

{chr(10).join(f"- `{s['id']}`（版本 {s['vintage'] or '—'}）" for s in sources)}
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def export(conn: sqlite3.Connection, cfg: Config) -> list[Path]:
    workbook = cfg.output["workbook"]
    registry = cfg.output["registry"]
    notes = cfg.output["notes"]
    return [
        export_workbook(conn, cfg, workbook),
        export_registry(conn, cfg, registry),
        export_notes(conn, cfg, notes),
    ]
